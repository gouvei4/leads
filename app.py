"""
Sistema de Prospeccao de Leads
Busca empresas via Google Places API, salva num banco SQLite local
e permite gerenciar status/observacoes por uma interface web.

Rodar:
    pip install -r requirements.txt
    python app.py
Depois abra http://localhost:5000 no navegador.
"""

import json
import os
import re
import sqlite3
import time
from datetime import datetime
from io import BytesIO

import requests
from flask import Flask, g, jsonify, request, send_file, render_template

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.path.join(BASE_DIR, "leads.db")
CONFIG_PATH = os.path.join(BASE_DIR, "config.json")

STATUS_OPTIONS = ["Novo", "Contatado", "Respondeu", "Negociando", "Cliente", "Sem interesse"]

app = Flask(__name__)


# ---------------------------------------------------------------- config ---

def load_config():
    if os.path.exists(CONFIG_PATH):
        with open(CONFIG_PATH, "r", encoding="utf-8") as f:
            return json.load(f)
    return {"api_key": os.environ.get("GOOGLE_PLACES_KEY", "")}


def save_config(cfg):
    with open(CONFIG_PATH, "w", encoding="utf-8") as f:
        json.dump(cfg, f, ensure_ascii=False, indent=2)


# ------------------------------------------------------------------ db -----

def get_db():
    if "db" not in g:
        g.db = sqlite3.connect(DB_PATH)
        g.db.row_factory = sqlite3.Row
    return g.db


@app.teardown_appcontext
def close_db(exception=None):
    db = g.pop("db", None)
    if db is not None:
        db.close()


def init_db():
    conn = sqlite3.connect(DB_PATH)
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS leads (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            nome TEXT NOT NULL,
            telefone TEXT,
            endereco TEXT,
            cidade TEXT,
            site TEXT,
            link_maps TEXT,
            termo_busca TEXT,
            status TEXT DEFAULT 'Novo',
            ultimo_contato TEXT,
            observacoes TEXT,
            criado_em TEXT,
            UNIQUE(nome, endereco)
        )
        """
    )
    conn.commit()
    conn.close()


# --------------------------------------------------------- google places ---

PLACES_URL = "https://places.googleapis.com/v1/places:searchText"
FIELD_MASK = ",".join([
    "places.displayName",
    "places.formattedAddress",
    "places.nationalPhoneNumber",
    "places.internationalPhoneNumber",
    "places.websiteUri",
    "places.googleMapsUri",
    "places.businessStatus",
    "nextPageToken",
])


def buscar_places(termo, cidade, api_key):
    """Busca um termo numa cidade, paginando ate ~60 resultados."""
    resultados = []
    body = {"textQuery": f"{termo} {cidade}", "languageCode": "pt-BR"}
    headers = {
        "Content-Type": "application/json",
        "X-Goog-Api-Key": api_key,
        "X-Goog-FieldMask": FIELD_MASK,
    }

    for _pagina in range(3):
        resp = requests.post(PLACES_URL, json=body, headers=headers, timeout=30)
        if resp.status_code != 200:
            raise RuntimeError(f"Google Places API erro {resp.status_code}: {resp.text[:300]}")

        data = resp.json()
        for place in data.get("places", []):
            if place.get("businessStatus") == "CLOSED_PERMANENTLY":
                continue
            resultados.append({
                "nome": place.get("displayName", {}).get("text", ""),
                "endereco": place.get("formattedAddress", ""),
                "telefone": place.get("nationalPhoneNumber") or place.get("internationalPhoneNumber", ""),
                "site": place.get("websiteUri", ""),
                "link_maps": place.get("googleMapsUri", ""),
                "cidade": cidade,
                "termo_busca": termo,
            })

        token = data.get("nextPageToken")
        if not token:
            break
        body["pageToken"] = token
        time.sleep(2)

    return resultados


# ------------------------------------------------------------- routes ------

@app.route("/")
def index():
    return render_template("index.html", status_options=STATUS_OPTIONS)


@app.route("/api/config", methods=["GET", "POST"])
def api_config():
    if request.method == "POST":
        data = request.get_json(force=True)
        cfg = load_config()
        cfg["api_key"] = data.get("api_key", "").strip()
        save_config(cfg)
        return jsonify({"ok": True})
    cfg = load_config()
    tem_chave = bool(cfg.get("api_key"))
    return jsonify({"tem_chave": tem_chave})


@app.route("/api/leads", methods=["GET"])
def api_leads_list():
    db = get_db()
    cidade = request.args.get("cidade", "").strip()
    status = request.args.get("status", "").strip()
    termo = request.args.get("termo", "").strip()
    busca = request.args.get("q", "").strip()

    query = "SELECT * FROM leads WHERE 1=1"
    params = []
    if cidade:
        query += " AND cidade = ?"
        params.append(cidade)
    if status:
        query += " AND status = ?"
        params.append(status)
    if termo:
        query += " AND termo_busca = ?"
        params.append(termo)
    if busca:
        query += " AND (nome LIKE ? OR endereco LIKE ?)"
        params.extend([f"%{busca}%", f"%{busca}%"])
    query += " ORDER BY cidade, nome"

    rows = db.execute(query, params).fetchall()
    leads = [dict(r) for r in rows]

    cidades = [r["cidade"] for r in db.execute("SELECT DISTINCT cidade FROM leads ORDER BY cidade").fetchall()]
    termos = [r["termo_busca"] for r in db.execute("SELECT DISTINCT termo_busca FROM leads ORDER BY termo_busca").fetchall()]

    return jsonify({"leads": leads, "cidades": cidades, "termos": termos, "status_options": STATUS_OPTIONS})


@app.route("/api/leads/<int:lead_id>", methods=["PATCH"])
def api_lead_update(lead_id):
    data = request.get_json(force=True)
    campos_permitidos = {"status", "ultimo_contato", "observacoes"}
    updates = {k: v for k, v in data.items() if k in campos_permitidos}
    if not updates:
        return jsonify({"ok": False, "erro": "nada para atualizar"}), 400

    db = get_db()
    set_clause = ", ".join(f"{k} = ?" for k in updates)
    params = list(updates.values()) + [lead_id]
    db.execute(f"UPDATE leads SET {set_clause} WHERE id = ?", params)
    db.commit()
    return jsonify({"ok": True})


@app.route("/api/leads/<int:lead_id>", methods=["DELETE"])
def api_lead_delete(lead_id):
    db = get_db()
    db.execute("DELETE FROM leads WHERE id = ?", (lead_id,))
    db.commit()
    return jsonify({"ok": True})


@app.route("/api/buscar", methods=["POST"])
def api_buscar():
    cfg = load_config()
    api_key = cfg.get("api_key", "")
    if not api_key:
        return jsonify({"ok": False, "erro": "Configure sua chave da Google Places API na aba Configuracoes."}), 400

    data = request.get_json(force=True)
    termos = [t.strip() for t in data.get("termos", []) if t.strip()]
    cidades = [c.strip() for c in data.get("cidades", []) if c.strip()]

    if not termos or not cidades:
        return jsonify({"ok": False, "erro": "Informe pelo menos 1 termo e 1 cidade."}), 400

    db = get_db()
    novos = 0
    duplicados = 0
    erros = []
    log = []

    for cidade in cidades:
        for termo in termos:
            try:
                resultados = buscar_places(termo, cidade, api_key)
            except Exception as exc:  # noqa: BLE001
                erros.append(f"{termo} / {cidade}: {exc}")
                log.append(f"ERRO em '{termo}' - {cidade}: {exc}")
                continue

            adicionados_nesta_busca = 0
            for r in resultados:
                try:
                    db.execute(
                        """
                        INSERT INTO leads (nome, telefone, endereco, cidade, site, link_maps, termo_busca, status, criado_em)
                        VALUES (?, ?, ?, ?, ?, ?, ?, 'Novo', ?)
                        """,
                        (
                            r["nome"], r["telefone"], r["endereco"], r["cidade"],
                            r["site"], r["link_maps"], r["termo_busca"],
                            datetime.now().isoformat(timespec="seconds"),
                        ),
                    )
                    novos += 1
                    adicionados_nesta_busca += 1
                except sqlite3.IntegrityError:
                    duplicados += 1
            db.commit()
            log.append(f"'{termo}' em {cidade}: {len(resultados)} encontrados, {adicionados_nesta_busca} novos")
            time.sleep(1)

    return jsonify({
        "ok": True,
        "novos": novos,
        "duplicados": duplicados,
        "erros": erros,
        "log": log,
    })


@app.route("/api/export")
def api_export():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    db = get_db()
    rows = db.execute("SELECT * FROM leads ORDER BY cidade, nome").fetchall()

    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"

    headers = ["Nome", "Telefone", "Cidade", "Endereco", "Site", "Link Maps",
               "Termo da Busca", "Status", "Ultimo Contato", "Observacoes"]
    ws.append(headers)

    header_fill = PatternFill("solid", start_color="1F2937", end_color="1F2937")
    header_font = Font(bold=True, color="FFFFFF")
    for col in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=col)
        c.fill = header_fill
        c.font = header_font

    for r in rows:
        ws.append([
            r["nome"], r["telefone"], r["cidade"], r["endereco"], r["site"],
            r["link_maps"], r["termo_busca"], r["status"], r["ultimo_contato"], r["observacoes"],
        ])

    widths = [32, 16, 18, 42, 26, 14, 18, 14, 14, 30]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:J{len(rows) + 1}"

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(
        buf,
        as_attachment=True,
        download_name=f"leads_{datetime.now().strftime('%Y%m%d')}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.route("/api/stats")
def api_stats():
    db = get_db()
    total = db.execute("SELECT COUNT(*) c FROM leads").fetchone()["c"]
    por_status = {
        r["status"]: r["c"]
        for r in db.execute("SELECT status, COUNT(*) c FROM leads GROUP BY status").fetchall()
    }
    por_cidade = [
        dict(r) for r in db.execute(
            "SELECT cidade, COUNT(*) c FROM leads GROUP BY cidade ORDER BY c DESC"
        ).fetchall()
    ]
    return jsonify({"total": total, "por_status": por_status, "por_cidade": por_cidade})


if __name__ == "__main__":
    init_db()
    print(f"Banco de dados em: {DB_PATH}")
    print("Abra http://localhost:5000 no navegador")
    app.run(debug=True, port=5000)
