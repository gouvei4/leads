"""
Importa leads de um arquivo CSV ou XLSX para o banco de dados do sistema,
sem duplicar quem ja existir (mesmo nome + endereco).

Se o arquivo tiver colunas de Status, Observacoes ou Ultimo Contato
(como a planilha CRM gerada anteriormente), esses valores sao
aproveitados em vez de resetar tudo para "Novo".

Uso:
    python import_leads.py leads_cacambas.csv
    python import_leads.py leads_cacambas_CRM.xlsx
"""
import sqlite3
import sys
from datetime import datetime

DB_PATH = "leads.db"

# nomes de coluna aceitos (minusculo, sem acento) -> campo interno
MAPA_COLUNAS = {
    "nome": "nome",
    "telefone": "telefone",
    "endereco": "endereco",
    "endereço": "endereco",
    "cidade": "cidade",
    "site": "site",
    "link_maps": "link_maps",
    "link maps": "link_maps",
    "termo_busca": "termo_busca",
    "termo da busca": "termo_busca",
    "busca": "busca_completa",  # formato antigo: "termo - cidade"
    "status": "status",
    "ultimo_contato": "ultimo_contato",
    "ultimo contato": "ultimo_contato",
    "último contato": "ultimo_contato",
    "observacoes": "observacoes",
    "observações": "observacoes",
}


def normalizar(texto):
    subs = str.maketrans("áàâãéêíóôõúüç", "aaaaeeiooouuc")
    return texto.strip().lower().translate(subs)


def ler_csv(caminho):
    import csv
    with open(caminho, "r", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        linhas = [dict(row) for row in reader]
        colunas_originais = reader.fieldnames or []
    return linhas, colunas_originais


def ler_xlsx(caminho):
    from openpyxl import load_workbook
    wb = load_workbook(caminho, data_only=True)
    # tenta achar uma aba chamada "Leads"; senao usa a primeira
    ws = wb["Leads"] if "Leads" in wb.sheetnames else wb[wb.sheetnames[0]]

    linhas_brutas = list(ws.iter_rows(values_only=True))
    if not linhas_brutas:
        return [], []
    colunas_originais = [str(c) if c is not None else "" for c in linhas_brutas[0]]
    linhas = []
    for row in linhas_brutas[1:]:
        d = {}
        for col_nome, valor in zip(colunas_originais, row):
            d[col_nome] = valor if valor is not None else ""
        linhas.append(d)
    return linhas, colunas_originais


def mapear_linha(linha_bruta, colunas_originais):
    """Converte uma linha (dict com nomes de coluna originais) pro formato interno."""
    resultado = {}
    for col_original in colunas_originais:
        chave_norm = normalizar(col_original)
        campo_interno = MAPA_COLUNAS.get(chave_norm)
        if not campo_interno:
            continue
        valor = linha_bruta.get(col_original, "")
        resultado[campo_interno] = str(valor).strip() if valor is not None else ""

    # formato antigo do CSV: coluna "busca" tipo "termo - cidade"
    if "busca_completa" in resultado and not resultado.get("termo_busca"):
        termo, _, cidade = resultado["busca_completa"].partition(" - ")
        resultado["termo_busca"] = termo.strip()
        if not resultado.get("cidade"):
            resultado["cidade"] = cidade.strip()

    # link "Abrir no Maps" -> a planilha CRM troca o texto do link por isso;
    # sem o valor real, deixamos vazio em vez de salvar o texto errado
    if resultado.get("link_maps", "").lower() in ("abrir no maps", ""):
        resultado["link_maps"] = ""

    return resultado


def main():
    if len(sys.argv) < 2:
        print("Uso: python import_leads.py caminho/do/arquivo.csv (ou .xlsx)")
        sys.exit(1)

    caminho = sys.argv[1]
    ext = caminho.rsplit(".", 1)[-1].lower()

    if ext == "csv":
        linhas, colunas = ler_csv(caminho)
    elif ext in ("xlsx", "xlsm"):
        linhas, colunas = ler_xlsx(caminho)
    else:
        print(f"Formato '.{ext}' nao suportado. Use .csv ou .xlsx.")
        sys.exit(1)

    if not linhas:
        print("Nenhuma linha encontrada no arquivo.")
        sys.exit(1)

    conn = sqlite3.connect(DB_PATH)
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS leads (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            nome TEXT NOT NULL,
            telefone TEXT,
            endereco TEXT,
            cidade TEXT,
            site TEXT,
            link_maps TEXT,
            termo_busca TEXT,
            status TEXT DEFAULT 'Novo',
            ultimo_contato TEXT,
            observacoes TEXT,
            criado_em TEXT,
            UNIQUE(nome, endereco)
        )
        """
    )

    novos = 0
    atualizados = 0
    ignorados = 0

    for linha_bruta in linhas:
        item = mapear_linha(linha_bruta, colunas)
        nome = item.get("nome", "").strip()
        if not nome:
            ignorados += 1
            continue

        try:
            conn.execute(
                """
                INSERT INTO leads (nome, telefone, endereco, cidade, site, link_maps,
                                    termo_busca, status, ultimo_contato, observacoes, criado_em)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    nome,
                    item.get("telefone", ""),
                    item.get("endereco", ""),
                    item.get("cidade", ""),
                    item.get("site", ""),
                    item.get("link_maps", ""),
                    item.get("termo_busca", ""),
                    item.get("status") or "Novo",
                    item.get("ultimo_contato", ""),
                    item.get("observacoes", ""),
                    datetime.now().isoformat(timespec="seconds"),
                ),
            )
            novos += 1
        except sqlite3.IntegrityError:
            # ja existe (mesmo nome+endereco) -> atualiza status/observacoes se vieram preenchidos
            campos_upd = []
            valores_upd = []
            for campo in ("status", "ultimo_contato", "observacoes"):
                if item.get(campo):
                    campos_upd.append(f"{campo} = ?")
                    valores_upd.append(item[campo])
            if campos_upd:
                valores_upd.extend([nome, item.get("endereco", "")])
                conn.execute(
                    f"UPDATE leads SET {', '.join(campos_upd)} WHERE nome = ? AND endereco = ?",
                    valores_upd,
                )
                atualizados += 1

    conn.commit()
    conn.close()
    print(f"Novos: {novos} | Ja existiam (atualizados): {atualizados} | Ignorados (sem nome): {ignorados}")


if __name__ == "__main__":
    main()
